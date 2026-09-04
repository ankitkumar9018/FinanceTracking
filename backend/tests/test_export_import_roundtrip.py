"""Every data export must be re-importable.

The exports write human-readable headers ("Stock Symbol", "Avg Price",
"Type") while the import templates use machine keys ("stock_symbol",
"price", "transaction_type"). These tests pin the header-alias layer that
bridges the two, plus the position-snapshot semantics for exports that carry
no transaction ledger (the holdings CSV / the "Holdings" sheet).

Every test is a *true* round-trip through the service functions — export,
parse, import into a fresh portfolio, compare cumulative quantity and average
price against the source. No network, no HTTP layer.

Run with:
    uv run --no-sync pytest tests/test_export_import_roundtrip.py -q
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.holding import Holding
from app.models.portfolio import Portfolio
from app.models.transaction import Transaction
from app.models.user import User
from app.services.csv_import_service import (
    canonical_column,
    generate_csv_template,
    parse_csv,
)
from app.services.excel_service import (
    _TEMPLATE_COLUMNS,
    export_portfolio,
    generate_template,
    import_to_portfolio,
    parse_excel,
)
from app.services.export_service import (
    export_holdings_csv,
    export_transactions_csv,
    export_workbook_xlsx,
)
from app.services.portfolio_service import calculate_cumulative_holding

# ---------------------------------------------------------------------------
# Seed helpers
# ---------------------------------------------------------------------------

_LEDGER = {
    "RELIANCE": [
        ("BUY", date(2024, 1, 15), 10.0, 2500.0, 50.0),
        ("BUY", date(2024, 2, 20), 5.0, 2600.0, 25.0),
        ("SELL", date(2024, 3, 10), 3.0, 2700.0, 30.0),
    ],
    "SAP": [
        ("BUY", date(2024, 4, 1), 20.0, 120.50, 5.0),
    ],
}

_EXCHANGES = {"RELIANCE": "NSE", "SAP": "XETRA"}


async def _seed_source(db: AsyncSession) -> Portfolio:
    """Create a user + portfolio with two holdings and a small ledger."""
    user = User(email="roundtrip@example.com", password_hash="x", display_name="RT")
    db.add(user)
    await db.flush()

    portfolio = Portfolio(user_id=user.id, name="Source", currency="INR")
    db.add(portfolio)
    await db.flush()

    for symbol, txns in _LEDGER.items():
        holding = Holding(
            portfolio_id=portfolio.id,
            stock_symbol=symbol,
            stock_name=f"{symbol} Ltd",
            exchange=_EXCHANGES[symbol],
            cumulative_quantity=0.0,
            average_price=0.0,
            current_price=2900.0 if symbol == "RELIANCE" else 130.0,
            current_rsi=55.5,
            sector="Energy" if symbol == "RELIANCE" else "Software",
            lower_mid_range_1=2300.0,
            upper_mid_range_1=2700.0,
            base_level=2000.0,
            top_level=3200.0,
            notes="seeded",
        )
        db.add(holding)
        await db.flush()
        for tx_type, tx_date, qty, price, brokerage in txns:
            db.add(
                Transaction(
                    holding_id=holding.id,
                    transaction_type=tx_type,
                    date=tx_date,
                    quantity=qty,
                    price=price,
                    brokerage=brokerage,
                    source="MANUAL",
                )
            )
        await db.flush()
        await calculate_cumulative_holding(holding.id, db)

    return portfolio


async def _fresh_portfolio(db: AsyncSession, name: str = "Target") -> Portfolio:
    """A second, empty portfolio owned by the same user."""
    user = (await db.execute(select(User))).scalars().first()
    if user is None:
        user = User(email="roundtrip@example.com", password_hash="x", display_name="RT")
        db.add(user)
        await db.flush()
    portfolio = Portfolio(user_id=user.id, name=name, currency="INR")
    db.add(portfolio)
    await db.flush()
    return portfolio


async def _positions(db: AsyncSession, portfolio_id: int) -> dict[str, tuple[float, float]]:
    """{symbol: (cumulative_quantity, average_price)} for a portfolio."""
    result = await db.execute(
        select(Holding).where(Holding.portfolio_id == portfolio_id)
    )
    return {
        h.stock_symbol: (float(h.cumulative_quantity), float(h.average_price))
        for h in result.scalars().all()
    }


def _assert_positions_match(
    source: dict[str, tuple[float, float]],
    imported: dict[str, tuple[float, float]],
) -> None:
    assert set(imported) == set(source), (source, imported)
    for symbol, (qty, avg) in source.items():
        got_qty, got_avg = imported[symbol]
        assert got_qty == pytest.approx(qty, rel=1e-6), symbol
        assert got_avg == pytest.approx(avg, rel=1e-6), symbol


# ---------------------------------------------------------------------------
# 1. Transactions CSV — the headline round-trip
# ---------------------------------------------------------------------------


async def test_transactions_csv_round_trips_the_whole_ledger(db: AsyncSession):
    source = await _seed_source(db)
    csv_text = await export_transactions_csv(source.id, db)

    # The export keeps its human-readable headers — that is a feature.
    assert csv_text.splitlines()[0] == (
        "Stock Symbol,Exchange,Type,Date,Quantity,Price,Brokerage,Source,Notes"
    )

    parsed = parse_csv(csv_text.encode())
    assert len(parsed) == 4  # every ledger row survived
    assert {r["transaction_type"] for r in parsed} == {"BUY", "SELL"}
    assert date(2024, 1, 15) in {r["date"] for r in parsed}

    target = await _fresh_portfolio(db)
    summary = await import_to_portfolio(parsed, target.id, db, source="CSV")
    assert summary["holdings_created"] == 2
    assert summary["transactions_created"] == 4
    assert summary["transactions_skipped"] == 0

    _assert_positions_match(
        await _positions(db, source.id), await _positions(db, target.id)
    )

    # Brokerage and dates survived too (the ledger, not just the totals).
    imported_tx = (
        await db.execute(
            select(Transaction)
            .join(Holding, Transaction.holding_id == Holding.id)
            .where(Holding.portfolio_id == target.id)
        )
    ).scalars().all()
    assert sorted(float(t.brokerage) for t in imported_tx) == [5.0, 25.0, 30.0, 50.0]
    assert all(t.source == "CSV" for t in imported_tx)


async def test_transactions_csv_reimport_is_deduped(db: AsyncSession):
    source = await _seed_source(db)
    csv_text = await export_transactions_csv(source.id, db)
    target = await _fresh_portfolio(db)

    first = await import_to_portfolio(parse_csv(csv_text.encode()), target.id, db)
    second = await import_to_portfolio(parse_csv(csv_text.encode()), target.id, db)

    assert first["transactions_created"] == 4
    assert second["transactions_created"] == 0
    assert second["transactions_skipped"] == 4
    _assert_positions_match(
        await _positions(db, source.id), await _positions(db, target.id)
    )


# ---------------------------------------------------------------------------
# 2. Holdings CSV — a position snapshot, not a ledger
# ---------------------------------------------------------------------------


async def test_holdings_csv_imports_as_position_snapshot(db: AsyncSession):
    source = await _seed_source(db)
    csv_text = await export_holdings_csv(source.id, db)
    # The six zone levels and notes trail the computed columns: they are user
    # input, so the export has to carry them or a CSV restore silently wipes
    # the 5-zone configuration (see test_stream_c_import_export.py).
    assert csv_text.splitlines()[0] == (
        "Stock Symbol,Stock Name,Exchange,Quantity,Avg Price,"
        "Current Price,P&L %,Action Needed,RSI,Sector,"
        "Lower Mid 1,Lower Mid 2,Upper Mid 1,Upper Mid 2,"
        "Base Level,Top Level,Notes"
    )

    parsed = parse_csv(csv_text.encode())
    assert len(parsed) == 2
    # No transaction type / date in the file → one synthetic opening BUY,
    # dated today, at the average price (NOT the current price).
    assert all(r["transaction_type"] == "BUY" for r in parsed)
    assert all(r["date"] == date.today() for r in parsed)
    reliance = next(r for r in parsed if r["stock_symbol"] == "RELIANCE")
    assert reliance["price"] == pytest.approx(2533.3333, rel=1e-6)  # avg, not 2900
    assert reliance["sector"] == "Energy"

    target = await _fresh_portfolio(db)
    summary = await import_to_portfolio(parsed, target.id, db, source="CSV")
    assert summary["holdings_created"] == 2
    assert summary["transactions_created"] == 2

    _assert_positions_match(
        await _positions(db, source.id), await _positions(db, target.id)
    )


async def test_holdings_csv_reimport_does_not_double_the_position(db: AsyncSession):
    source = await _seed_source(db)
    csv_text = await export_holdings_csv(source.id, db)
    target = await _fresh_portfolio(db)

    before = await _positions(db, source.id)
    first = await import_to_portfolio(parse_csv(csv_text.encode()), target.id, db)
    second = await import_to_portfolio(parse_csv(csv_text.encode()), target.id, db)

    assert first["transactions_created"] == 2
    assert second["transactions_created"] == 0
    assert second["transactions_skipped"] == 2
    _assert_positions_match(before, await _positions(db, target.id))


async def test_zero_quantity_snapshot_row_is_skipped(db: AsyncSession):
    csv_text = (
        "Stock Symbol,Stock Name,Exchange,Quantity,Avg Price,"
        "Current Price,P&L %,Action Needed,RSI,Sector\n"
        "RELIANCE,Reliance Ltd,NSE,10,2500.0,2600.0,4.0,N,55,Energy\n"
        "EXITED,Exited Co,NSE,0,0,,,N,,Energy\n"
    )
    parsed = parse_csv(csv_text.encode())
    assert [r["stock_symbol"] for r in parsed] == ["RELIANCE"]


async def test_unknown_columns_are_ignored(db: AsyncSession):
    csv_text = (
        "Stock Symbol,Exchange,Type,Date,Quantity,Price,Brokerage,Source,Notes,"
        "Broker Ref,Whatever\n"
        "INFY,NSE,BUY,2024-05-02,7,1500.0,10.0,MANUAL,note,XYZ-1,42\n"
    )
    parsed = parse_csv(csv_text.encode())
    assert len(parsed) == 1
    assert parsed[0]["quantity"] == 7.0
    assert parsed[0]["price"] == 1500.0
    # The transactions export omits the display name — fall back to the symbol.
    assert parsed[0]["stock_name"] == "INFY"

    target = await _fresh_portfolio(db, name="Extras")
    summary = await import_to_portfolio(parsed, target.id, db, source="CSV")
    assert summary["transactions_created"] == 1


def test_partial_transaction_row_is_still_rejected():
    """A type without a date (or vice-versa) is malformed, not a snapshot."""
    header = "Stock Symbol,Stock Name,Exchange,Type,Date,Quantity,Price\n"
    assert parse_csv((header + "INFY,Infosys,NSE,BUY,,7,1500\n").encode()) == []
    assert parse_csv((header + "INFY,Infosys,NSE,,2024-05-02,7,1500\n").encode()) == []


# ---------------------------------------------------------------------------
# 3. Excel export → import
# ---------------------------------------------------------------------------


async def test_excel_export_round_trips_via_transactions_sheet(db: AsyncSession):
    source = await _seed_source(db)
    xlsx = await export_portfolio(source.id, db)

    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Holdings", "Transactions"]
    wb.close()

    parsed = parse_excel(xlsx)
    # The Transactions sheet is preferred over the Holdings snapshot: four
    # ledger rows with their real dates, not two rows dated today.
    assert len(parsed) == 4
    assert {r["date"] for r in parsed} == {
        date(2024, 1, 15), date(2024, 2, 20), date(2024, 3, 10), date(2024, 4, 1),
    }
    assert {r["transaction_type"] for r in parsed} == {"BUY", "SELL"}

    target = await _fresh_portfolio(db)
    summary = await import_to_portfolio(parsed, target.id, db, source="EXCEL")
    assert summary["holdings_created"] == 2
    assert summary["transactions_created"] == 4

    _assert_positions_match(
        await _positions(db, source.id), await _positions(db, target.id)
    )


async def test_excel_holdings_sheet_is_the_fallback(db: AsyncSession):
    """Without a Transactions sheet, the Holdings snapshot must still import."""
    source = await _seed_source(db)
    wb = load_workbook(io.BytesIO(await export_portfolio(source.id, db)))
    del wb["Transactions"]
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    parsed = parse_excel(buf.getvalue())
    assert len(parsed) == 2
    assert all(r["transaction_type"] == "BUY" for r in parsed)
    assert all(r["date"] == date.today() for r in parsed)
    reliance = next(r for r in parsed if r["stock_symbol"] == "RELIANCE")
    assert reliance["price"] == pytest.approx(2533.3333, rel=1e-6)
    # "Lower Mid 1" etc. alias back onto the template's range columns.
    assert reliance["lower_mid_range_1"] == pytest.approx(2300.0)
    assert reliance["top_level"] == pytest.approx(3200.0)

    target = await _fresh_portfolio(db)
    await import_to_portfolio(parsed, target.id, db, source="EXCEL")
    _assert_positions_match(
        await _positions(db, source.id), await _positions(db, target.id)
    )


async def test_multi_sheet_workbook_falls_back_to_holdings(db: AsyncSession):
    """``export_workbook_xlsx``'s Transactions sheet has no Exchange column, so
    the parser must fall back to its Holdings snapshot instead of failing."""
    source = await _seed_source(db)
    parsed = parse_excel(await export_workbook_xlsx(source.id, db))

    assert len(parsed) == 2
    assert all(r["date"] == date.today() for r in parsed)
    target = await _fresh_portfolio(db)
    await import_to_portfolio(parsed, target.id, db, source="EXCEL")
    _assert_positions_match(
        await _positions(db, source.id), await _positions(db, target.id)
    )


# ---------------------------------------------------------------------------
# 4. Regression — the templates import exactly as before
# ---------------------------------------------------------------------------


def test_csv_template_still_imports_unchanged():
    parsed = parse_csv(generate_csv_template().encode())
    assert len(parsed) == 1
    row = parsed[0]
    assert row["stock_symbol"] == "RELIANCE"
    assert row["stock_name"] == "Reliance Industries Ltd"
    assert row["exchange"] == "NSE"
    assert row["transaction_type"] == "BUY"
    assert row["date"] == date(2024, 1, 15)
    assert row["quantity"] == 10.0
    assert row["price"] == 2500.0
    assert row["brokerage"] == 50.0
    assert row["lower_mid_range_1"] == 2300.0
    assert row["lower_mid_range_2"] == 2100.0
    assert row["upper_mid_range_1"] == 2700.0
    assert row["upper_mid_range_2"] == 2900.0
    assert row["base_level"] == 2000.0
    assert row["top_level"] == 3000.0
    assert row["sector"] == "Energy"
    assert row["notes"] == "Initial purchase"


def test_excel_template_still_imports_unchanged():
    parsed = parse_excel(generate_template())
    assert len(parsed) == 1
    row = parsed[0]
    assert row["stock_symbol"] == "RELIANCE"
    assert row["stock_name"] == "Reliance Industries Ltd"
    assert row["transaction_type"] == "BUY"
    assert row["date"] == date(2024, 1, 15)
    assert row["quantity"] == 10.0
    assert row["price"] == 2500.0
    assert row["sector"] == "Energy"
    assert row["notes"] == "Initial purchase"


def test_every_template_header_maps_to_itself():
    """Zero-regression guarantee for the alias layer."""
    for column in _TEMPLATE_COLUMNS:
        assert canonical_column(column) == column


def test_export_headers_map_onto_template_keys():
    assert canonical_column("Stock Symbol") == "stock_symbol"
    assert canonical_column("Stock Name") == "stock_name"
    assert canonical_column("Type") == "transaction_type"
    assert canonical_column("Date") == "date"
    assert canonical_column("Avg Price") == "price"
    assert canonical_column("Quantity") == "quantity"
    assert canonical_column("Brokerage") == "brokerage"
    assert canonical_column("Fees") == "brokerage"
    assert canonical_column("Sector") == "sector"
    assert canonical_column("Notes") == "notes"
    # Export-only columns stay out of the way (must NOT become `price`,
    # and "Action" must NOT become a transaction type).
    assert canonical_column("Current Price") == "current_price"
    assert canonical_column("P&L %") == "p_l"
    assert canonical_column("Action Needed") == "action_needed"
    assert canonical_column("Action") == "action"
    assert canonical_column("RSI") == "rsi"


def test_single_sheet_upload_without_named_sheets_still_parses():
    """A third-party sheet named anything at all still goes through the
    active-sheet path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Symbol", "Name", "Exchange", "Type", "Date", "Qty", "Price", "Fees"])
    ws.append(["INFY", "Infosys", "NSE", "buy", "2024-05-02", 7, 1500.0, 10.0])
    buf = io.BytesIO()
    wb.save(buf)

    parsed = parse_excel(buf.getvalue())
    assert len(parsed) == 1
    assert parsed[0]["stock_symbol"] == "INFY"
    assert parsed[0]["transaction_type"] == "BUY"
    assert parsed[0]["quantity"] == 7.0
    assert parsed[0]["brokerage"] == 10.0
