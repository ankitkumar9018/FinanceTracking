"""Import/export regression tests (stream C).

Each class pins one defect that was found in the import/export surface:

1. ``TestBackupKeepsFundTypeAndCustomFields`` — the JSON "full backup" dropped
   ``Holding.fund_type`` and ``Holding.custom_fields``, so restoring your own
   backup turned every German equity ETF from 30% partially exempt into 0%
   (overstating tax on every later sale) and wiped stop-loss / target-allocation.
2. ``TestMalformedUploadsAre400`` — malformed CSV/JSON uploads raised unhandled
   exceptions (HTTP 500) instead of an actionable 400.
3. ``TestBankStatementRowsAreNotHoldings`` — OFX/QIF bank lines were imported as
   one fake holding per payee, adding every credit to Total Invested.
4. ``TestSkippedRowsAreReported`` — the endpoints reported only the rows that
   survived parsing, so a mostly-failed import still rendered "Import Successful".
5. ``TestHoldingsCsvKeepsZoneConfig`` — the holdings CSV/XLSX exported the
   computed action zone but none of the six levels it is derived from, so a
   CSV backup/restore lost the whole 5-zone configuration.

Run with:
    cd backend && uv run pytest tests/test_stream_c_import_export.py -q
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.holding import Holding
from app.models.portfolio import Portfolio
from app.models.transaction import Transaction
from app.models.user import User
from app.services.backup_service import export_portfolio_json, import_portfolio_json
from app.services.csv_import_service import parse_csv
from app.services.excel_service import import_to_portfolio
from app.services.export_service import export_holdings_csv, export_workbook_xlsx
from app.services.ofx_qif_import_service import (
    import_statement,
    parse_ofx,
    parse_qif,
    split_cash_rows,
)
from app.services.tax_service import teilfreistellung_for_fund_type

# ---------------------------------------------------------------------------
# Seeding helpers
# ---------------------------------------------------------------------------


async def _make_user(db: AsyncSession, email: str) -> User:
    user = User(email=email, password_hash="x", display_name="Stream C")
    db.add(user)
    await db.flush()
    return user


async def _make_portfolio(
    db: AsyncSession, user: User, name: str = "Main", currency: str = "EUR"
) -> Portfolio:
    portfolio = Portfolio(user_id=user.id, name=name, currency=currency)
    db.add(portfolio)
    await db.flush()
    return portfolio


async def _seed_etf_portfolio(db: AsyncSession, email: str) -> Portfolio:
    """A German equity ETF with a Teilfreistellung class, zone levels, notes and
    a populated ``custom_fields`` (stop-loss + target allocation)."""
    user = await _make_user(db, email)
    portfolio = await _make_portfolio(db, user)
    holding = Holding(
        portfolio_id=portfolio.id,
        stock_symbol="EUNL",
        stock_name="iShares Core MSCI World",
        exchange="XETRA",
        currency="EUR",
        fund_type="EQUITY_ETF",
        custom_fields={"stop_loss_price": 68.5, "target_allocation_pct": 12.5},
        cumulative_quantity=0.0,
        average_price=0.0,
        lower_mid_range_1=70.0,
        lower_mid_range_2=75.0,
        upper_mid_range_1=95.0,
        upper_mid_range_2=100.0,
        base_level=60.0,
        top_level=110.0,
        sector="Diversified",
        notes="Core position — do not trim below 60",
    )
    db.add(holding)
    await db.flush()
    db.add(
        Transaction(
            holding_id=holding.id,
            transaction_type="BUY",
            date=date(2024, 3, 1),
            quantity=100.0,
            price=80.0,
            brokerage=1.0,
            source="MANUAL",
        )
    )
    await db.flush()
    holding.cumulative_quantity = 100.0
    holding.average_price = 80.0
    await db.flush()
    return portfolio


async def _holdings_of(db: AsyncSession, portfolio_id: int) -> list[Holding]:
    rows = await db.execute(
        select(Holding).where(Holding.portfolio_id == portfolio_id)
    )
    return list(rows.scalars().all())


async def _create_portfolio_via_api(
    client: AsyncClient, auth_headers: dict[str, str], name: str = "Imports"
) -> int:
    resp = await client.post(
        "/api/v1/portfolios/",
        json={"name": name, "currency": "INR"},
        headers=auth_headers,
    )
    assert resp.status_code == 201, resp.text
    return int(resp.json()["id"])


# ===========================================================================
# 1. JSON backup must carry fund_type and custom_fields
# ===========================================================================


class TestBackupKeepsFundTypeAndCustomFields:
    async def test_export_serialises_both_columns(self, db: AsyncSession):
        portfolio = await _seed_etf_portfolio(db, "backup-export@example.com")

        backup = await export_portfolio_json(portfolio.id, portfolio.user_id, db)

        holding = backup["portfolio"]["holdings"][0]
        assert holding["fund_type"] == "EQUITY_ETF"
        assert holding["custom_fields"] == {
            "stop_loss_price": 68.5,
            "target_allocation_pct": 12.5,
        }

    async def test_restore_keeps_teilfreistellung_and_stop_loss(
        self, db: AsyncSession
    ):
        """Restoring must not silently re-tax the position at 0% exemption."""
        portfolio = await _seed_etf_portfolio(db, "backup-restore@example.com")
        backup = await export_portfolio_json(portfolio.id, portfolio.user_id, db)

        summary = await import_portfolio_json(backup, portfolio.user_id, db)

        restored = await _holdings_of(db, summary["portfolio_id"])
        assert len(restored) == 1
        assert restored[0].fund_type == "EQUITY_ETF"
        # The money consequence: 30% of the gain stays exempt, as before.
        assert teilfreistellung_for_fund_type(restored[0].fund_type) == 30.0
        assert restored[0].custom_fields == {
            "stop_loss_price": 68.5,
            "target_allocation_pct": 12.5,
        }

    async def test_restore_defaults_missing_and_malformed_custom_fields(
        self, db: AsyncSession
    ):
        """An older backup (no keys) and a hand-edited one (wrong type) both
        restore to an empty dict, never to ``None`` or a scalar."""
        user = await _make_user(db, "backup-legacy@example.com")
        backup = {
            "format": "financetracker_backup",
            "version": "1.0",
            "portfolio": {
                "name": "Legacy",
                "currency": "EUR",
                "holdings": [
                    {
                        "stock_symbol": "OLD",
                        "stock_name": "Old Format",
                        "exchange": "XETRA",
                        "transactions": [],
                        "dividends": [],
                    },
                    {
                        "stock_symbol": "BAD",
                        "stock_name": "Bad Custom Fields",
                        "exchange": "XETRA",
                        "custom_fields": "not-an-object",
                        "transactions": [],
                        "dividends": [],
                    },
                ],
            },
        }

        summary = await import_portfolio_json(backup, user.id, db)

        restored = {h.stock_symbol: h for h in await _holdings_of(db, summary["portfolio_id"])}
        assert restored["OLD"].fund_type is None
        assert restored["OLD"].custom_fields == {}
        assert restored["BAD"].custom_fields == {}


# ===========================================================================
# 2. Malformed uploads → 400, not 500
# ===========================================================================


_OVERLONG_CSV = (
    "stock_symbol,stock_name,exchange,transaction_type,date,quantity,price\n"
    'INFY,"' + "a" * 200_000 + '",NSE,BUY,2024-01-15,10,1500\n'
).encode()


class TestMalformedUploadsAre400:
    @pytest.mark.parametrize(
        "path_template",
        [
            "/api/v1/import-export/csv?portfolio_id={pid}",
            "/api/v1/import-export/csv/dividends?portfolio_id={pid}",
            "/api/v1/import-export/csv/mutual-funds?portfolio_id={pid}",
        ],
    )
    async def test_overlong_field_is_rejected_with_400(
        self,
        client: AsyncClient,
        auth_headers: dict[str, str],
        path_template: str,
    ):
        pid = await _create_portfolio_via_api(client, auth_headers)

        resp = await client.post(
            path_template.format(pid=pid),
            files={"file": ("broker.csv", _OVERLONG_CSV, "text/csv")},
            headers=auth_headers,
        )

        assert resp.status_code == 400, resp.text
        assert "parse" in resp.json()["detail"].lower()

    async def test_overlong_field_on_tax_records_is_rejected_with_400(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        resp = await client.post(
            "/api/v1/import-export/csv/tax-records",
            files={"file": ("tax.csv", _OVERLONG_CSV, "text/csv")},
            headers=auth_headers,
        )

        assert resp.status_code == 400, resp.text
        assert "parse" in resp.json()["detail"].lower()

    async def test_backup_missing_required_field_is_400(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        """A truncated backup (transaction without a price) is a bad request."""
        truncated = (
            b'{"format": "financetracker_backup", "version": "1.0", "portfolio": '
            b'{"name": "P", "currency": "INR", "holdings": [{"stock_symbol": "INFY", '
            b'"stock_name": "Infosys", "exchange": "NSE", "transactions": '
            b'[{"transaction_type": "BUY", "date": "2024-01-15", "quantity": 10}]}]}}'
        )

        resp = await client.post(
            "/api/v1/import-export/json",
            files={"file": ("backup.json", truncated, "application/json")},
            headers=auth_headers,
        )

        assert resp.status_code == 400, resp.text
        assert "price" in resp.json()["detail"]

    async def test_backup_with_wrong_field_type_is_400(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        wrong_type = (
            b'{"format": "financetracker_backup", "version": "1.0", "portfolio": '
            b'{"name": "P", "currency": "INR", "holdings": [{"stock_symbol": "INFY", '
            b'"stock_name": "Infosys", "exchange": "NSE", "transactions": '
            b'[{"transaction_type": "BUY", "date": "2024-01-15", "quantity": null, '
            b'"price": 1500}]}]}}'
        )

        resp = await client.post(
            "/api/v1/import-export/json",
            files={"file": ("backup.json", wrong_type, "application/json")},
            headers=auth_headers,
        )

        assert resp.status_code == 400, resp.text
        assert "Malformed backup file" in resp.json()["detail"]

    async def test_json_array_backup_is_400(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        resp = await client.post(
            "/api/v1/import-export/json",
            files={"file": ("backup.json", b"[1, 2, 3]", "application/json")},
            headers=auth_headers,
        )

        assert resp.status_code == 400, resp.text

    async def test_a_rejected_restore_leaves_no_partial_portfolio(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        """The half-built portfolio must be rolled back, not left behind."""
        before = await client.get("/api/v1/portfolios/", headers=auth_headers)
        truncated = (
            b'{"format": "financetracker_backup", "version": "1.0", "portfolio": '
            b'{"name": "Half", "currency": "INR", "holdings": [{"stock_symbol": "INFY", '
            b'"stock_name": "Infosys", "exchange": "NSE", "transactions": '
            b'[{"transaction_type": "BUY", "date": "2024-01-15", "quantity": 10}]}]}}'
        )

        resp = await client.post(
            "/api/v1/import-export/json",
            files={"file": ("backup.json", truncated, "application/json")},
            headers=auth_headers,
        )
        assert resp.status_code == 400

        after = await client.get("/api/v1/portfolios/", headers=auth_headers)
        assert len(after.json()) == len(before.json())
        assert all("Half" not in p["name"] for p in after.json())


# ===========================================================================
# 3. Bank/cash statement lines are not holdings
# ===========================================================================


OFX_BANK_ONLY = """OFXHEADER:100
DATA:OFXSGML
VERSION:102

<OFX>
<BANKMSGSRSV1>
<STMTTRNRS>
<STMTRS>
<BANKTRANLIST>
<STMTTRN>
<TRNTYPE>CREDIT
<DTPOSTED>20240131120000
<TRNAMT>4500.00
<NAME>ACME PAYROLL
<MEMO>Salary January
</STMTTRN>
<STMTTRN>
<TRNTYPE>DEBIT
<DTPOSTED>20240201120000
<TRNAMT>-120.55
<NAME>SUPERMARKET
<MEMO>Groceries
</STMTTRN>
</BANKTRANLIST>
</STMTRS>
</STMTTRNRS>
</BANKMSGSRSV1>
</OFX>
"""

QIF_MIXED = """!Type:Bank
D01/31/2024
T4500.00
PACME PAYROLL
MSalary January
^
!Type:Invest
D02/15/2024
NBuy
YAcme Corp
Q10
I100.00
T1000.00
^
"""


class TestBankStatementRowsAreNotHoldings:
    def test_bank_rows_are_tagged_cash_and_split_out(self):
        rows = parse_ofx(OFX_BANK_ONLY.encode())
        assert len(rows) == 2  # still parsed, so we can explain them

        investment, cash = split_cash_rows(rows)
        assert investment == []
        assert [r["stock_symbol"] for r in cash] == ["ACME PAYROLL", "SUPERMARKET"]

    async def test_import_statement_creates_no_holding_for_cash(
        self, db: AsyncSession
    ):
        user = await _make_user(db, "ofx-bank@example.com")
        portfolio = await _make_portfolio(db, user, name="Bank", currency="INR")

        summary = await import_statement(
            parse_ofx(OFX_BANK_ONLY.encode()), portfolio.id, db, source="OFX"
        )

        assert summary["holdings_created"] == 0
        assert summary["transactions_created"] == 0
        assert summary["cash_rows_skipped"] == 2
        assert await _holdings_of(db, portfolio.id) == []

    async def test_bank_only_ofx_upload_is_rejected(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        pid = await _create_portfolio_via_api(client, auth_headers, name="Bank OFX")

        resp = await client.post(
            f"/api/v1/import-export/import/ofx?portfolio_id={pid}",
            files={"file": ("statement.ofx", OFX_BANK_ONLY.encode(), "text/plain")},
            headers=auth_headers,
        )

        assert resp.status_code == 400, resp.text
        assert "bank/cash" in resp.json()["detail"]

        holdings = await client.get(
            f"/api/v1/holdings/?portfolio_id={pid}", headers=auth_headers
        )
        assert holdings.json() == []

    async def test_mixed_qif_imports_trades_and_reports_skipped_cash(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        pid = await _create_portfolio_via_api(client, auth_headers, name="Mixed QIF")

        resp = await client.post(
            f"/api/v1/import-export/import/qif?portfolio_id={pid}",
            files={"file": ("statement.qif", QIF_MIXED.encode(), "text/plain")},
            headers=auth_headers,
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["rows_read"] == 2
        assert body["rows_parsed"] == 1
        assert body["cash_rows_skipped"] == 1
        assert body["holdings_created"] == 1
        assert body["transactions_created"] == 1
        assert "bank/cash" in body["warning"]

        holdings = await client.get(
            f"/api/v1/holdings/?portfolio_id={pid}", headers=auth_headers
        )
        symbols = [h["stock_symbol"] for h in holdings.json()]
        assert symbols == ["ACMECORP"]
        assert "ACME PAYROLL" not in symbols

    def test_qif_bank_record_is_tagged_cash(self):
        investment, cash = split_cash_rows(parse_qif(QIF_MIXED.encode()))
        assert [r["stock_symbol"] for r in investment] == ["ACMECORP"]
        assert [r["stock_symbol"] for r in cash] == ["ACME PAYROLL"]


# ===========================================================================
# 4. Rows the parser dropped are reported
# ===========================================================================


class TestSkippedRowsAreReported:
    async def test_partial_csv_import_reports_the_dropped_rows(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        pid = await _create_portfolio_via_api(client, auth_headers, name="Partial")
        csv_bytes = (
            b"stock_symbol,stock_name,exchange,transaction_type,date,quantity,price\n"
            b"RELIANCE,Reliance Industries,NSE,BUY,2024-01-15,10,2500\n"
            b"INFY,Infosys,NSE,BUY,not-a-date,5,1500\n"
            b"TCS,TCS,NSE,BUY,2024-01-16,3,\n"
            b",,,,,,\n"
        )

        resp = await client.post(
            f"/api/v1/import-export/csv?portfolio_id={pid}",
            files={"file": ("trades.csv", csv_bytes, "text/csv")},
            headers=auth_headers,
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        # The blank trailing line carries no data and is not counted as a row.
        assert body["rows_read"] == 3
        assert body["rows_parsed"] == 1
        assert body["rows_skipped"] == 2
        assert "2 of 3" in body["warning"]
        assert body["transactions_created"] == 1

    async def test_clean_csv_import_reports_no_skips_and_no_warning(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        pid = await _create_portfolio_via_api(client, auth_headers, name="Clean")
        csv_bytes = (
            b"stock_symbol,stock_name,exchange,transaction_type,date,quantity,price\n"
            b"RELIANCE,Reliance Industries,NSE,BUY,2024-01-15,10,2500\n"
            b"INFY,Infosys,NSE,BUY,2024-01-16,5,1500\n"
        )

        resp = await client.post(
            f"/api/v1/import-export/csv?portfolio_id={pid}",
            files={"file": ("trades.csv", csv_bytes, "text/csv")},
            headers=auth_headers,
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["rows_read"] == 2
        assert body["rows_parsed"] == 2
        assert body["rows_skipped"] == 0
        assert "warning" not in body

    async def test_partial_dividend_import_reports_the_dropped_rows(
        self, client: AsyncClient, auth_headers: dict[str, str]
    ):
        pid = await _create_portfolio_via_api(client, auth_headers, name="Divs")
        holdings_csv = (
            b"stock_symbol,stock_name,exchange,transaction_type,date,quantity,price\n"
            b"RELIANCE,Reliance Industries,NSE,BUY,2024-01-15,10,2500\n"
        )
        seed = await client.post(
            f"/api/v1/import-export/csv?portfolio_id={pid}",
            files={"file": ("trades.csv", holdings_csv, "text/csv")},
            headers=auth_headers,
        )
        assert seed.status_code == 200, seed.text

        dividends_csv = (
            b"stock_symbol,exchange,ex_date,payment_date,amount_per_share,total_amount\n"
            b"RELIANCE,NSE,2024-02-01,2024-02-15,8,80\n"
            b"RELIANCE,NSE,nonsense,2024-03-15,8,80\n"
        )
        resp = await client.post(
            f"/api/v1/import-export/csv/dividends?portfolio_id={pid}",
            files={"file": ("divs.csv", dividends_csv, "text/csv")},
            headers=auth_headers,
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["rows_read"] == 2
        assert body["rows_parsed"] == 1
        assert body["rows_skipped"] == 1
        assert "dividend rows" in body["warning"]


# ===========================================================================
# 5. Holdings CSV / XLSX keep the zone configuration
# ===========================================================================


class TestHoldingsCsvKeepsZoneConfig:
    async def test_csv_export_round_trips_zones_and_notes(self, db: AsyncSession):
        source = await _seed_etf_portfolio(db, "zones-csv@example.com")
        csv_text = await export_holdings_csv(source.id, db)

        header = csv_text.splitlines()[0]
        for column in (
            "Lower Mid 1", "Lower Mid 2", "Upper Mid 1", "Upper Mid 2",
            "Base Level", "Top Level", "Notes",
        ):
            assert column in header

        parsed = parse_csv(csv_text.encode())
        assert len(parsed) == 1
        assert parsed[0]["lower_mid_range_1"] == 70.0
        assert parsed[0]["top_level"] == 110.0

        user = await _make_user(db, "zones-csv-target@example.com")
        target = await _make_portfolio(db, user, name="Restored")
        await import_to_portfolio(parsed, target.id, db, source="CSV")

        restored = await _holdings_of(db, target.id)
        assert len(restored) == 1
        assert float(restored[0].lower_mid_range_1) == 70.0
        assert float(restored[0].lower_mid_range_2) == 75.0
        assert float(restored[0].upper_mid_range_1) == 95.0
        assert float(restored[0].upper_mid_range_2) == 100.0
        assert float(restored[0].base_level) == 60.0
        assert float(restored[0].top_level) == 110.0
        assert restored[0].notes == "Core position — do not trim below 60"

    async def test_csv_export_leaves_unset_zones_empty(self, db: AsyncSession):
        user = await _make_user(db, "zones-empty@example.com")
        portfolio = await _make_portfolio(db, user, name="No Zones", currency="INR")
        db.add(
            Holding(
                portfolio_id=portfolio.id,
                stock_symbol="INFY",
                stock_name="Infosys",
                exchange="NSE",
                currency="INR",
                cumulative_quantity=10.0,
                average_price=1500.0,
            )
        )
        await db.flush()

        csv_text = await export_holdings_csv(portfolio.id, db)
        data_row = csv_text.splitlines()[1]

        # Six empty zone cells + an empty notes cell, not "None".
        assert "None" not in data_row
        assert data_row.endswith(",,,,,,,")

    async def test_notes_cannot_smuggle_a_formula_into_the_new_column(
        self, db: AsyncSession
    ):
        """The new Notes column is free text — it must not become a live formula."""
        user = await _make_user(db, "zones-formula@example.com")
        portfolio = await _make_portfolio(db, user, name="Injection", currency="INR")
        db.add(
            Holding(
                portfolio_id=portfolio.id,
                stock_symbol="INFY",
                stock_name="Infosys",
                exchange="NSE",
                currency="INR",
                cumulative_quantity=10.0,
                average_price=1500.0,
                notes="=1+1",
            )
        )
        await db.flush()

        csv_text = await export_holdings_csv(portfolio.id, db)
        assert csv_text.splitlines()[1].endswith("'=1+1")

        wb = load_workbook(io.BytesIO(await export_workbook_xlsx(portfolio.id, db)))
        notes_cell = wb["Holdings"].cell(row=2, column=16)
        assert notes_cell.value == "'=1+1"
        assert notes_cell.data_type == "s"  # string, not a formula

    async def test_xlsx_holdings_sheet_carries_zone_columns(self, db: AsyncSession):
        source = await _seed_etf_portfolio(db, "zones-xlsx@example.com")

        wb = load_workbook(io.BytesIO(await export_workbook_xlsx(source.id, db)))
        ws = wb["Holdings"]
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]

        assert header[:2] == ["Symbol", "Exchange"]  # existing columns unmoved
        zones = dict(zip(header, row))
        assert zones["Lower Mid 1"] == 70.0
        assert zones["Top Level"] == 110.0
        assert zones["Notes"] == "Core position — do not trim below 60"
